"""
One-off script, run once and commit the resulting .docx templates:

  1. Adds a small "Signature:" label above each signature line (both the
     company's -- above the embedded signature image -- and the
     candidate's -- above their blank line), so it's clear what that line
     is for.
  2. The company's "Date:" line (under "For Sloud Solutions") now reads
     "Date: [DD Month YYYY]", reusing the existing letterDate placeholder
     that offer-letter.ts already fills in everywhere it appears -- so the
     company's signing date is auto-filled with the letter's date at
     generation time. The candidate's "Date:" line is left blank; that's
     filled in when they actually sign.
"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

TEMPLATES = [
    "offerletter/full-time/sloud-fulltime-letter-template.docx",
    "offerletter/part-time/sloud-parttime-letter-template.docx",
    "offerletter/internship/sloud-internship-letter-template.docx",
]

FOR_SLOUD_TEXT = "For Sloud Solutions"
ACCEPTED_PREFIXES = ("Accepted by", "Accepted")
COMPANY_DATE_OLD = "Date: ____________________"
COMPANY_DATE_NEW = "Date: [DD Month YYYY]"
LABEL_TEXT = "Signature:"


def iter_body_paragraphs(doc):
    for child in doc.element.body.iterchildren():
        yield from _iter_element_paragraphs(child, doc)


def _iter_element_paragraphs(element, doc):
    from docx.text.paragraph import Paragraph
    from docx.table import Table

    if element.tag == qn("w:p"):
        yield Paragraph(element, doc)
    elif element.tag == qn("w:tbl"):
        table = Table(element, doc)
        for row in table.rows:
            for cell in row.cells:
                for child in cell._tc.iterchildren():
                    yield from _iter_element_paragraphs(child, doc)


def add_label_before(paragraph):
    """Insert a small gray 'Signature:' label paragraph right before the given one."""
    label = paragraph.insert_paragraph_before()
    label.paragraph_format.space_after = Pt(1.5)
    run = label.add_run(LABEL_TEXT)
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)


def process(path: str):
    doc = Document(path)
    paragraphs = list(iter_body_paragraphs(doc))

    for_idx = next(i for i, p in enumerate(paragraphs) if p.text.strip() == FOR_SLOUD_TEXT)
    accepted_idx = next(
        i for i, p in enumerate(paragraphs) if p.text.strip().startswith(ACCEPTED_PREFIXES)
    )

    # Company: label above the signature image/line, and auto-filled date two paragraphs later.
    add_label_before(paragraphs[for_idx + 1])
    date_para = paragraphs[for_idx + 3]
    assert date_para.text.strip() == COMPANY_DATE_OLD.strip(), f"unexpected company date text: {date_para.text!r}"
    for run in list(date_para.runs):
        run.text = ""
        run._element.getparent().remove(run._element)
    new_run = date_para.add_run(COMPANY_DATE_NEW)
    new_run.font.name = "Calibri"
    new_run.font.size = Pt(9.5)
    new_run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # Candidate: label above their blank signature line. Date stays blank -- filled when they sign.
    # Re-fetch paragraphs since the doc tree changed above.
    paragraphs = list(iter_body_paragraphs(doc))
    accepted_idx = next(
        i for i, p in enumerate(paragraphs) if p.text.strip().startswith(ACCEPTED_PREFIXES)
    )
    add_label_before(paragraphs[accepted_idx + 1])

    doc.save(path)
    print("updated", path)


for t in TEMPLATES:
    process(t)
