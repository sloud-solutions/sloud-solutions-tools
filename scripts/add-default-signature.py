"""
One-off script to attach the company's default signature image to the
"For Sloud Solutions" signature line in each offer-letter .docx template.

Replaces the blank "________________________________" line directly under
"For Sloud Solutions" with a small, cropped signature image (kept at a
modest ~0.8in height, aspect-locked). The candidate's own signature line
("Accepted by the Employee") is left untouched -- that one is signed by
the candidate, not pre-filled by us.

Run once locally (`python3 scripts/add-default-signature.py`) and commit
the resulting .docx files -- these are binary templates, not something the
app generates at build/runtime.
"""
from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn

TEMPLATES = [
    "offerletter/full-time/sloud-fulltime-letter-template.docx",
    "offerletter/part-time/sloud-parttime-letter-template.docx",
    "offerletter/internship/sloud-internship-letter-template.docx",
]
SIGNATURE_IMAGE = "public/images/brand/signature-default.png"
SIGNATURE_HEIGHT_IN = 0.8

FOR_SLOUD_TEXT = "For Sloud Solutions"
BLANK_LINE_TEXT = "________________________________"


def iter_body_paragraphs(doc):
    """Yield every paragraph in document order, including inside tables (any depth)."""
    parent_elm = doc.element.body
    for child in parent_elm.iterchildren():
        yield from _iter_element_paragraphs(child, doc)


def _iter_element_paragraphs(element, doc):
    from docx.text.paragraph import Paragraph
    from docx.table import Table

    tag = element.tag
    if tag == qn("w:p"):
        yield Paragraph(element, doc)
    elif tag == qn("w:tbl"):
        table = Table(element, doc)
        for row in table.rows:
            for cell in row.cells:
                for child in cell._tc.iterchildren():
                    yield from _iter_element_paragraphs(child, doc)


def add_signature(path: str):
    doc = Document(path)
    paragraphs = list(iter_body_paragraphs(doc))

    for_idx = next(i for i, p in enumerate(paragraphs) if p.text.strip() == FOR_SLOUD_TEXT)
    blank_idx = next(
        i for i in range(for_idx + 1, len(paragraphs)) if paragraphs[i].text.strip() == BLANK_LINE_TEXT
    )
    target = paragraphs[blank_idx]

    for run in list(target.runs):
        run.text = ""
        run._element.getparent().remove(run._element)

    run = target.add_run()
    run.add_picture(SIGNATURE_IMAGE, height=Inches(SIGNATURE_HEIGHT_IN))

    doc.save(path)
    print("updated", path)


for t in TEMPLATES:
    add_signature(t)
